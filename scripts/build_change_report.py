# -*- coding: utf-8 -*-
"""Sinh bao cao thay doi (tieng Viet) tu cac vung boi vang trong main_APIN_highlight.tex."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "BaoCaoNgay02072026.docx"

YELLOW = RGBColor(0xBF, 0x8F, 0x00)
DARK = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# --- base style ---
base = doc.styles["Normal"]
base.font.name = "Times New Roman"
base.font.size = Pt(12)

def set_heading(p, size, color=DARK, space_before=10, space_after=4):
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = color
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)

def h1(text):
    p = doc.add_paragraph(text)
    set_heading(p, 15, DARK, 14, 6)
    return p

def h2(text):
    p = doc.add_paragraph(text)
    set_heading(p, 13, YELLOW, 10, 3)
    return p

def para(text, italic=False, color=None, size=12, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.font.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)

# ============================ TITLE ============================
t = doc.add_paragraph("BÁO CÁO CHỈNH SỬA BÀI BÁO APIN")
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.runs[0].font.size = Pt(18); t.runs[0].font.bold = True; t.runs[0].font.color.rgb = DARK
sub = doc.add_paragraph("Tổng hợp các thay đổi so với phiên bản trước (theo phần bôi vàng)")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(12.5); sub.runs[0].font.italic = True; sub.runs[0].font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
rm = meta.add_run("Ngày 02/07/2026  •  Tệp: main_APIN.tex (bản sạch) và main_APIN_highlight.tex (bản bôi vàng)")
rm.font.size = Pt(11); rm.font.color.rgb = GREY
doc.add_paragraph()

# ============================ 1. TONG QUAN ============================
h1("1. Định hướng chung của đợt sửa")
para(
    "Đợt chỉnh sửa lần này không thêm mô hình mới, mà thay đổi cách diễn giải và định vị "
    "kết quả nhằm trả lời hai nhóm nhận xét của hội đồng: (i) làm rõ giá trị ứng dụng của "
    "công trình và (ii) củng cố tính chặt chẽ của phần bằng chứng. Trục xuyên suốt là "
    "chuyển từ thông điệp cũ — xem rò rỉ đồ thị như một vấn đề tái lập với kết quả “không "
    "đáng kể” — sang một khung lập luận mới gọi là “tác hại có điều kiện”.")
para("Bốn thay đổi lớn nhất về nội dung:")
bullets([
    ("Khung “tác hại có điều kiện” (mô hình hai nhân tố). ",
     "Rò rỉ đồ thị chỉ làm dịch chỉ số AUC khi đồng thời cao ở hai yếu tố: lưu lượng rò rỉ "
     "(phần bằng chứng của một cạnh đến từ dữ liệu bị giữ lại để kiểm thử) và mức phụ thuộc "
     "của mô hình nền vào đồ thị. Nếu một trong hai yếu tố xấp xỉ 0 thì AUC gần như đứng yên, "
     "nên AUC là một “chuông báo” không đáng tin về rò rỉ."),
    ("Định vị lại quy trình như một lớp hỗ trợ ra quyết định và chấm điểm rủi ro. ",
     "Thay vì mô tả công trình là “hạ tầng tái lập”, bài nhấn mạnh giá trị vận hành: quy trình "
     "giúp người triển khai biết một tập dữ liệu đang ở chế độ rủi ro nào trước khi độ chính "
     "xác kịp bộc lộ điều đó."),
    ("Bổ sung phép kiểm thử huấn luyện lại ở hạ nguồn có mỏ neo (đối chứng dương). ",
     "Đây là bằng chứng mới cho thấy kết luận về DDR là có điều kiện: chỉ đúng trên mô hình nền "
     "thực sự “đọc” đồ thị."),
    ("Tăng tính minh bạch và khả năng đọc. ",
     "Thêm mã giả cho ba thuật toán, các ví dụ minh hoạ bằng số, một bảng tương tác hai chiều, "
     "một bảng “từ tín hiệu rà soát đến hành động triển khai”, cùng các câu tự giới hạn phạm vi "
     "kết luận (guardrail)."),
])

# ============================ 2. CHI TIET THEO PHAN ============================
h1("2. Chi tiết thay đổi theo từng phần")

h2("2.1. Tóm tắt (Abstract)")
para("Viết lại đoạn kết theo khung hai nhân tố và định vị mới. Cụ thể:")
bullets([
    "Nêu rõ rò rỉ chỉ dịch AUC khi lưu lượng rò rỉ và mức phụ thuộc đồ thị cùng cao.",
    "Dẫn số liệu đối lập: bộ lọc chặt giữ lưu lượng thấp nên biến thể gộp toàn bộ nhật ký chỉ "
    "làm AUC đổi tối đa 0,003; nhưng khi chủ động tăng lưu lượng bằng tiêm nhiễm, các mô hình "
    "phụ thuộc đồ thị dịch mạnh (GKT khoảng +0,05; GIKT khoảng +0,03), còn mô hình chỉ dùng "
    "chuỗi thì gần như không đổi.",
    "Khẳng định AUC “im lặng” ở ba trong bốn chế độ, nên quy trình đo trực tiếp lưu lượng "
    "(khối lượng cạnh dựng, chỉ số TBMR) thay vì dựa vào độ chính xác.",
    "Chốt định vị: quy trình là lớp hỗ trợ ra quyết định và chấm điểm rủi ro rò rỉ, không phải "
    "con đường để tăng độ chính xác công bố.",
])

h2("2.2. Phần mở đầu (Introduction)")
bullets([
    ("Câu luận điểm trung tâm. ",
     "Thêm đoạn nêu rõ rò rỉ đồ thị là “có hại một cách có điều kiện” chứ không nguy hiểm đồng "
     "loạt, và giải thích vì sao cần một lớp rà soát đo lưu lượng độc lập với độ chính xác."),
    ("Hai ví dụ minh hoạ cụ thể. ",
     "Một ví dụ về cạnh rò rỉ (một quan hệ tiền đề có bằng chứng tăng từ 40 lên 60 nhờ 20 lượt "
     "của người học bị giữ lại) để định nghĩa trực quan “lưu lượng”; và một ví dụ về khái niệm "
     "kiến thức khởi đầu nguội (cold-start)."),
    ("Đoạn “người thực hành làm gì khác đi”. ",
     "Liệt kê bốn quyết định cụ thể trước khi công bố/triển khai: dựng đồ thị sạch theo fold, "
     "đọc chỉ báo lưu lượng thay vì độ chính xác, ràng buộc mọi kết luận hạ nguồn bằng kiểm tra "
     "đối chứng, và ưu tiên phép tăng cường dữ liệu bảo toàn cấu trúc tiền đề."),
    ("Đổi cách định vị. ",
     "Thay cụm mô tả cũ bằng “lớp rà soát hỗ trợ ra quyết định và chấm điểm rủi ro”."),
])

h2("2.3. Danh mục đóng góp (mục C2)")
para(
    "Viết lại phần bằng chứng hạ nguồn: DDR dự báo được mức sụt độ chính xác trên mô hình nền "
    "phụ thuộc đồ thị đã vượt qua kiểm tra đối chứng (GKT trên XES3G5M, hệ số tương quan Pearson "
    "r = 0,97; toán tử bảo toàn tiền đề gây mất AUC ít nhất ở cùng ngân sách nhiễu), nhưng chỉ là "
    "rà soát cấu trúc thuần tuý trên các mô hình nền trơ với đồ thị. Bổ sung câu tự giới hạn: "
    "không tuyên bố DDR dự báo được độ chính xác cho các mô hình không dùng đồ thị.")

h2("2.4. Mã giả cho ba thuật toán")
para("Bổ sung đoạn giới thiệu và tiêu đề cho ba thuật toán được trình bày dưới dạng mã giả:")
bullets([
    "Thuật toán rà soát rò rỉ theo nguồn gốc cạnh cho từng fold (gồm cờ cấu trúc và các chỉ số lưu lượng).",
    "Thuật toán dựng đồ thị tiền đề chỉ từ dữ liệu huấn luyện kèm kiểm tra tính phi chu trình.",
    "Thuật toán tính DAG Disruption Rate (DDR).",
])

h2("2.5. §4.3 — Thí nghiệm tiêm nhiễm cạnh có kiểm soát")
bullets([
    ("Nâng vai trò thí nghiệm. ",
     "Câu mở đầu được viết lại để đóng khung đây là “sự hiện thực hoá có kiểm soát” của ô nguy "
     "hiểm trong bảng hai chiều, thay vì chỉ là phép kiểm tra độ nhạy của chỉ số."),
    ("Bổ sung GIKT vào kết luận. ",
     "Nêu rõ cả hai mô hình phụ thuộc đồ thị cùng dịch (GKT 0,810→0,860; GIKT 0,852→0,880) trong "
     "khi mô hình chỉ dùng chuỗi đứng yên quanh 0,850–0,858."),
    ("Thêm bảng tương tác hai chiều. ",
     "Bảng thể hiện AUC chỉ dịch ở ô “lưu lượng cao × phụ thuộc cao”; ba ô còn lại xấp xỉ 0, và "
     "điểm vận hành thực tế nằm ở hàng lưu lượng thấp."),
    ("Ví dụ minh hoạ theo số. ",
     "Diễn giải hàng tiêm nhiễm 20%: khối lượng cạnh dựng tăng từ 1162 lên 1417 (+21,9%), TBMR "
     "tăng từ 0,754 lên 0,779, trong khi cờ cấu trúc vẫn bằng 0 — tức rà soát phát hiện trước khi "
     "AUC kịp phản ứng."),
    ("Đoạn giải thích sau bảng. ",
     "Cùng một mô hình GKT cho kết quả xấp xỉ 0 khi gộp toàn bộ nhật ký nhưng +0,05 khi tiêm "
     "nhiễm, chứng minh lưu lượng là biến điều khiển thật; tác hại có tính nhân, và trục phụ thuộc "
     "được kiểm chứng độc lập qua kiểm tra đối chứng."),
])

h2("2.6. §4.5 — Khử đồ thị (dựng chỉ từ huấn luyện so với gộp toàn bộ nhật ký)")
para(
    "Đọc lại kết quả “không dịch” (≤0,003) như một phép đo chế độ vận hành, chứ không phải bằng "
    "chứng cho thấy rò rỉ vô hại. Do bộ lọc hỗ trợ chặt, phần lớn cạnh sống sót khi gộp toàn bộ "
    "vốn cũng đã sống sót khi chỉ dựng từ huấn luyện, nên lưu lượng thấp. Chính bộ lọc bóp nghẹt "
    "kênh rò rỉ ở đây, chứ không phải bản chất chung của rò rỉ; cùng kênh đó tải khoảng 0,05 AUC "
    "khi lưu lượng được nâng lên bằng tiêm nhiễm.")

h2("2.7. §4.6 — DDR như một “mục tiêu thiết kế” (có điều kiện)")
para(
    "Bổ sung câu làm rõ tính điều kiện: trên mô hình nền phụ thuộc đồ thị, phần thưởng cấu trúc "
    "của toán tử bảo toàn tiền đề còn kéo theo lợi thế ở hạ nguồn (mất AUC ít nhất ở cùng ngân "
    "sách trên GKT/XES3G5M), nên “mục tiêu thiết kế” chỉ được xác nhận về độ chính xác khi mô "
    "hình nền thực sự dùng đồ thị.")

h2("2.8. §4.7 — Kiểm thử hạ nguồn có mỏ neo (viết lại toàn bộ)")
para("Đây là phần được làm mới nhiều nhất, cấu trúc lại thành bốn ý:")
bullets([
    ("Nguyên tắc kiểm tra đối chứng. ",
     "Vì quan hệ DDR→AUC bằng 0 là mập mờ (có thể do DDR không dự báo, hoặc do mô hình bỏ qua đồ "
     "thị), mọi diễn giải đều được ràng buộc bằng một mỏ neo phá gần như toàn bộ đồ thị (p = 0,90)."),
    ("Mỏ neo phụ thuộc thấp. ",
     "DGEKT (mọi nơi) và GKT trên ASSISTments: phá đồ thị gần như hoàn toàn nhưng AUC gần như "
     "không đổi (≤0,002) — được đọc là mỏ neo phụ thuộc ≈ 0, không phải kết luận về DDR."),
    ("Mô hình nền phụ thuộc đồ thị. ",
     "GKT trên XES3G5M vượt kiểm tra đối chứng rõ rệt (phá gần hết đồ thị làm AUC giảm 0,07–0,09); "
     "trong vùng diễn giải được, DDR bám sát mức sụt AUC (Pearson r = 0,97) và thứ tự toán tử "
     "“bảo toàn tiền đề < xoá cạnh < xoá nút” giữ đúng ở mọi ngân sách nhiễu."),
    ("Ý nghĩa đối chứng dương. ",
     "Đây chính là đối chứng dương mà phép quét chỉ-DGEKT trước đây còn thiếu; DDR trở thành mục "
     "tiêu thiết kế đúng trong chế độ phụ thuộc cao mà kiểm tra đối chứng nhận diện."),
])

h2("2.9. §5.1 — Thảo luận (kết quả có điều kiện)")
para(
    "Bổ sung đoạn nêu phép kiểm thử hạ nguồn có mỏ neo làm sắc bén hoá thành một kết quả có điều "
    "kiện: nơi kiểm tra đối chứng xác nhận mô hình đọc đồ thị thì DDR dự báo mức sụt độ chính xác; "
    "nơi kiểm tra thất bại thì DDR chỉ còn là rà soát cấu trúc/tái lập. Vì vậy, các công trình tăng "
    "cường dữ liệu sau này nên báo cáo cả mức phá cấu trúc (DDR) lẫn một phép kiểm thử hạ nguồn có "
    "ràng buộc đối chứng.")

h2("2.10. §5.2 — Liên hệ rò rỉ ↔ khởi đầu nguội và bảng bản đồ quyết định")
bullets([
    ("Rủi ro rò rỉ tập trung ở khái niệm khởi đầu nguội. ",
     "Đoạn mới nối hai chẩn đoán vốn tách rời: lưu lượng rò rỉ lớn nhất ở nơi vùng lân cận trong "
     "tập huấn luyện thưa nhất; một khái niệm rất nguội chỉ với vài lượt huấn luyện có thể bị một "
     "cạnh rò rỉ chi phối. Do đó nên báo cáo nguồn gốc theo từng tầng, không chỉ tổng hợp."),
    ("Từ tín hiệu rà soát đến hành động triển khai. ",
     "Thêm một bảng “bản đồ quyết định”: mỗi tín hiệu rà soát ứng với một cách đọc và một hành "
     "động triển khai cụ thể, độc lập với AUC — chính là dạng vận hành của quy trình bốn bước ở "
     "phần mở đầu."),
])

h2("2.11. §5.3 — Giới hạn (guardrail) và hướng phát triển")
bullets([
    ("Phạm vi của khẳng định hai nhân tố. ",
     "Nêu thẳng hai giới hạn: (1) ô lưu lượng cao mới chỉ đạt được bằng tiêm nhiễm có kiểm soát, "
     "chưa chứng minh xảy ra tự nhiên; (2) trục phụ thuộc hiện dựa trên GKT và GIKT (dương) so với "
     "simpleKT và DGEKT (≈ 0), nên tương tác mới ở mức minh hoạ, chưa phải đặc tả đầy đủ."),
    ("Cập nhật độ phủ của DDR hạ nguồn. ",
     "Ghi rõ phép quét hồi quy lại đã bao phủ DGEKT và GKT (GKT kèm mỏ neo đối chứng), với "
     "GKT/XES3G5M báo cáo cho seed 42 và việc nhân bản đa-seed đang tiến hành."),
    ("Hướng phát triển. ",
     "Hoàn tất nhân bản đa-seed cho GKT, thêm một mô hình nền phụ thuộc đồ thị thứ hai ở cả hai "
     "mức lưu lượng, và ghép tương quan DDR ở mức cạnh với biến thể phá vỡ khả năng truy cập."),
])

h2("2.12. Kết luận")
para("Đồng bộ ba câu với khung mới:")
bullets([
    "Định vị quy trình là lớp hỗ trợ ra quyết định và chấm điểm rủi ro cho triển khai đáng tin cậy.",
    "Phép kiểm thử hạ nguồn có mỏ neo khiến cách đọc dự báo trở thành có điều kiện theo mức phụ thuộc.",
    "Các kết quả hội tụ dưới cách đọc hai nhân tố: AUC chỉ dịch khi lưu lượng và mức phụ thuộc cùng "
    "cao, và khi đó nó “thổi phồng” chứ không phân biệt được đâu là cải thiện thật.",
])

# ============================ 3. GHI CHU KY THUAT ============================
h1("3. Ghi chú kỹ thuật và sản phẩm bàn giao")
bullets([
    ("Hai bản song song. ",
     "main_APIN.tex là bản sạch (không dấu vết bôi vàng, dùng để nộp); main_APIN_highlight.tex là "
     "bản bôi vàng toàn bộ thay đổi để tiện đối chiếu. Cả hai đều biên dịch ra 57 trang, không lỗi."),
    ("Cơ chế bôi vàng. ",
     "Dùng công tắc bật/tắt: bôi vàng từng cụm chữ trong dòng và bôi vàng cả khối đoạn/thuật toán, "
     "xử lý an toàn với công thức toán và trích dẫn."),
    ("Sửa một lỗi biên dịch có sẵn. ",
     "Kiểu định dạng thư mục tài liệu tham khảo khiến một địa chỉ liên kết có dấu gạch dưới gây lỗi; "
     "đã bổ sung một dòng làm lệnh hiển thị liên kết an toàn với dấu gạch dưới. Đây là sửa kỹ thuật, "
     "không liên quan nội dung học thuật."),
])

# summary table
h2("Bảng tổng hợp số vùng bôi vàng")
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Light Grid Accent 1"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
hdr[0].paragraphs[0].add_run("Loại thay đổi được bôi vàng").font.bold = True
hdr[1].paragraphs[0].add_run("Số vùng").font.bold = True
rows = [
    ("Sửa/thêm trong dòng (cụm chữ, câu, chú thích bảng)", "20"),
    ("Khối đoạn/thuật toán viết mới hoặc viết lại", "12"),
    ("Tổng cộng", "32"),
]
for a, b in rows:
    c = tbl.add_row().cells
    c[0].paragraphs[0].add_run(a)
    c[1].paragraphs[0].add_run(b)

doc.add_paragraph()
note = doc.add_paragraph()
rn = note.add_run("Ghi chú: bản báo cáo này bám sát đúng các vùng đã bôi vàng trong main_APIN_highlight.tex; "
                  "các tên mô hình (GKT, GIKT, DGEKT, simpleKT), chỉ số (AUC, TBMR, DDR) và bộ dữ liệu "
                  "(XES3G5M, ASSISTments) được giữ nguyên theo thông lệ chuyên ngành.")
rn.font.italic = True; rn.font.size = Pt(10.5); rn.font.color.rgb = GREY

doc.save(str(OUT))
print(f"Saved {OUT}")
