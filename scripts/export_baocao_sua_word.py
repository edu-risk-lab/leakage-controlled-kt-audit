#!/usr/bin/env python3
"""Export revision change report to Word."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = val
    doc.add_paragraph()


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "paper" / "baocaosua17062026.docx"

    doc = Document()
    title = doc.add_heading(
        "Báo cáo thay đổi manuscript APIN (revision)", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run("Dự án: ").bold = True
    meta.add_run("Leakage-Controlled Concept Graph Construction and Cold-Start Diagnostic Protocol for KT\n")
    meta.add_run("File chính: ").bold = True
    meta.add_run("paper/main_APIN.tex, paper/supplementary.tex\n")
    meta.add_run("Baseline so sánh: ").bold = True
    meta.add_run("commit 501e18c6 (trước revision reviewer)\n")
    meta.add_run("Ngày báo cáo: ").bold = True
    meta.add_run("17/06/2026\n")
    meta.add_run("Commits chính: ").bold = True
    meta.add_run("a8429b3a (multi-seed GKT), 43aa5c11 (reviewer fixes + simpleKT30 cache)\n")

    doc.add_paragraph(
        "Tài liệu này tóm tắt các thay đổi nội dung so với phiên bản trước revision, "
        "bao gồm chỉnh sửa đã commit và một số chỉnh sửa local (caption Table 12, bản highlight vàng)."
    )

    add_heading(doc, "1. GKT ablation và C5 (M1 — epoch / multi-seed)", 1)
    add_table(
        doc,
        ["Trước", "Sau"],
        [
            [
                "“Epoch-matched” GKT 30ep, batch 64, 1 seed (42)",
                "“GKT epoch-extended” 30ep, batch 32; simpleKT giữ 10ep → chỉ partially matched",
            ],
            [
                "Pooled Δ CI [−0.048, −0.027]",
                "Pooled 9-fold Δ ≈ −0.038, CI [−0.040, −0.035] (seeds 17/42/1234)",
            ],
            ["Chỉ Table S21", "Tables S21–S22 (seed 42 + multi-seed replication)"],
        ],
    )
    add_bullets(
        doc,
        [
            "Phạm vi: Abstract, Intro, contributions C5, §4.2 baseline, Discussion, Limitations, Conclusion, Data availability.",
            "Macro mới: GKTthirtyseedFortydeltaci; input gkt_epoch_ablation_macros.tex.",
        ],
    )

    add_heading(doc, "2. SKT / DyGKT vs bảng chính (M2)", 1)
    add_bullets(
        doc,
        [
            "Thống nhất: SKT và DyGKT là fork wiring sanity checks — có trong Table 7–8 nhưng loại khỏi primary ranking claims.",
            "ACC/NLL chi tiết ở Supplementary Tables S1–S5.",
        ],
    )

    add_heading(doc, "3. Training parity & implementation (M4)", 1)
    add_bullets(
        doc,
        [
            "§4.2 bổ sung hyperparams cụ thể (lr, batch, epochs, emb_size GIKT, pyKT vs native).",
            "Ghi rõ không claim bit-identical parity giữa codebase.",
            "Nêu bug label-index GIKT đã sửa trước khi báo cáo Table CV.",
            "GKT ablation: GKT-only extension, graph rebuild theo seed; không full compute parity.",
        ],
    )

    add_heading(doc, "4. Related Work & Eq.(1) (M3 / M5)", 1)
    add_bullets(
        doc,
        [
            "Đoạn mới: Positioning against reproducibility and graph-audit practice — so sánh pyKT, label-leakage audits; 6 lớp audit protocol.",
            "Eq.(1): đổi từ “Objective” → “Design intent (not a deployed solver)” — pipeline greedy, λ=0, không joint optimiser.",
        ],
    )

    add_heading(doc, "5. Controlled injection §4.3 (M2)", 1)
    add_bullets(
        doc,
        [
            "Giải thích GKT 0.810→0.860 @ 20% injection: densify E_pre, graph-channel enrichment — không phải split hygiene.",
            "Nhấn mạnh AUC alone không đáng tin làm leakage alarm khi ECR_flag = 0.",
        ],
    )

    add_heading(doc, "6. DDR downstream / C2 (M4)", 1)
    add_bullets(
        doc,
        [
            "Table 11 (DDR→AUC): exploratory, DGEKT-only; DDR = structural diagnostic, không proxy AUC universal.",
            "Future work: ưu tiên GKT downstream sweep (playbook docs/DDR_DOWNSTREAM_GKT.md).",
        ],
    )

    add_heading(doc, "7. Phạm vi inferential & vai trò corpus", 1)
    add_bullets(
        doc,
        [
            "XES3G5M: primary cho C5; ASSIST/Junyi: secondary / saturated sanity.",
            "Wilcoxon 3-fold: nhấn Δ magnitude + S16 + 9-fold GKT ablation.",
            "Residual risk C5: bỏ “single seed 42” cho epoch ablation (đã có 3 seed).",
        ],
    )

    add_heading(doc, "8. Supplementary material", 1)
    add_bullets(
        doc,
        [
            "§S21: GKT epoch-extended (batch 32, simpleKT reference 10ep).",
            "§S22 mới: multi-seed (17/42/1234), bảng per-seed + pooled 9-fold.",
            "Index/abstract S và prose S16 cập nhật pooled Δ.",
        ],
    )

    add_heading(doc, "9. Table 12 / cold-start (caption khớp pipeline)", 1)
    add_table(
        doc,
        ["Trước", "Sau"],
        [
            [
                "“fold 0 test interactions per stratum”",
                "“fold 0 validation+test held-out interactions per KC-frequency stratum (strata from train-fold counts only; simpleKT)”",
            ],
        ],
    )
    add_bullets(
        doc,
        [
            "Đồng bộ: Figure 7, Table S12, scripts generate_phase_c_tables.py, generate_paper_artifacts.py.",
        ],
    )

    add_heading(doc, "10. Bản highlight revision (local)", 1)
    add_bullets(
        doc,
        [
            "main_APIN_version2.pdf: ~26 khối \\rev{} (nền vàng, mdframed) trên diff vs 501e18c6.",
            "Script: scripts/apply_revision_highlight.py.",
            "Lưu ý: gỡ \\rev{} và mdframed trước nộp bản final.",
        ],
    )

    add_heading(doc, "11. Hạng mục chưa hoàn thành / tạm đóng", 1)
    add_table(
        doc,
        ["Hạng mục", "Trạng thái"],
        [
            ["M7 (GT θ sweep, thêm fold)", "Tạm đóng — framing + optional θ sweep"],
            ["M4 GPU DDR downstream GKT", "Playbook sẵn, chưa chạy GPU"],
            ["simpleKT 30ep reference matched", "Cache ingest; AUC ≈ trio — chưa full parity narrative"],
        ],
    )

    add_heading(doc, "12. Tóm tắt một đoạn (cover letter)", 1)
    doc.add_paragraph(
        "Revision làm rõ GKT epoch-extended + 9-fold multi-seed (partially matched caps), "
        "SKT/DyGKT diagnostic-only, Eq.(1) design intent, injection mechanism, pyKT positioning, "
        "DGEKT-only DDR downstream, sửa caption n = valid+test held-out (Table 12), và bổ sung Tables S21–S22; "
        "contribution định vị audit/reproducibility, không claim leakage control tăng AUC lớn "
        "(≤0.003 full-log ablation trên ba benchmark công khai)."
    )

    add_heading(doc, "Phụ lục: Phản hồi reviewer (tóm tắt)", 1)
    add_heading(doc, "A. Motivation leakage vs Table 9 (≤0.003)", 2)
    add_bullets(
        doc,
        [
            "Contribution = audit/reproducibility, không phải “sửa leakage → tăng AUC”.",
            "Full-log ablation ≠ injection: Table 9 hỏi penalty metric dưới filter graph hiện tại; injection là stress test tiered audit.",
            "Leakage nguy hiểm = provenance/validity of comparison, không chỉ ΔAUC.",
        ],
    )
    add_heading(doc, "B. Table 11 DDR→AUC trên DGEKT (C2)", 2)
    add_bullets(
        doc,
        [
            "C2-structural: DDR phân biệt operator (giữ claim).",
            "C2-predictive: null trên DGEKT — không claim DDR dự báo AUC.",
            "GKT injection §4.3 cho thấy backbone graph-sensitive; cần GKT downstream cho C2-predictive.",
        ],
    )

    doc.add_paragraph()
    foot = doc.add_paragraph("— Hết báo cáo —")
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in foot.runs:
        run.font.size = Pt(10)

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
