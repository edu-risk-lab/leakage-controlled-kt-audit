"""Regenerate paper/submission_APIN/APIN_Review_Report.docx for internal GS review."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "submission_APIN" / "APIN_Review_Report.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell


def build() -> None:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = title.add_run("Applied Intelligence (APIN)\nInternal Pre-Submission Review Report")
    r.bold = True
    r.font.size = Pt(14)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub.add_run(
        "Simulated external reviewer perspective · For PGS Nguyễn Văn Hậu (internal QA before advisor sign-off)\n"
        "Not an official journal review"
    ).italic = True

    meta = [
        "Manuscript: Leakage-Controlled Concept Graph Construction and Cold-Start Diagnostic Protocol for Knowledge Tracing",
        "Authors: T. Dao Minh, K.-T. Nguyen, D. Nguyen Tien, Q. K. Ngo, V.-H. Nguyen, L. Hoang Son",
        "File reviewed: paper/submission_APIN/main_APIN.pdf (60 pages total; main body §1–§6: 34 pages; appendix from p.35)",
        "LaTeX: main_APIN.tex (2,630 lines) + auto-generated tables; git main @ b15f15a7 (25 July 2026)",
        "PDF SHA-256: 35b771c0327a791bae3c341ba10ab683cd9429d42dc0e49f28097f2a31ced398",
        "Review date: 25 July 2026 (updated after P0 revision sprint)",
        "Simulated recommendation (post-revision): Accept with minor optional polish / Ready for advisor review",
    ]
    for line in meta:
        add_para(doc, line)

    add_heading(doc, "0. Tóm tắt cho PGS (tiếng Việt)", 1)
    add_para(
        doc,
        "Đây là bản mô phỏng phản biện APIN nội bộ, dùng để chuẩn bị trước khi gửi thầy nghiệm thu. "
        "Sau đợt chỉnh theo nhận xét PGS (22/07/2026) và rà soát kỹ thuật (crossref 12/12, A3.5 GPU, C8 rút trang), "
        "bản thảo đã xử lý xong các mục chặn A1–A3, B4–B7, C8–C10 trong REVISION_PLAN. "
        "Không còn lỗi số liệu nội tại đã biết; gói gửi GS: paper/gui_GS_package_APIN_20260725.zip.",
    )
    add_bullet(doc, "Điểm mạnh giữ nguyên: định vị trung thực, DDR manipulation check, injection decoupling, Junyi GT báo cáo bất lợi minh bạch.")
    add_bullet(doc, "Điểm cần thầy xác nhận: A2 (−0.041 GKT vs simpleKT) — đã disclosure đầy đủ; có cần rerun compute-matched GPU không?")
    add_bullet(doc, "Điểm mô phỏng PB ngoài phạm vi revision: mở rộng related work (M3), generalization ngoài KT (M4), thêm folds (M2) — tuỳ chọn trước nộp EM.")
    add_bullet(doc, "Đề xuất: chấp nhận gửi GS/nộp với framing hiện tại; GPU compute-matched chỉ chạy nếu thầy yêu cầu cứng.")

    add_heading(doc, "A. Trạng thái bản thảo & bằng chứng", 1)
    add_table(
        doc,
        ["Hạng mục", "Trạng thái", "Bằng chứng"],
        [
            ["Crossref số liệu", "12/12 pass", "audit/crossref_auc_report.md (rerun 25/07)"],
            ["A1 injection S18", "Verified GPU", "§4.3 decoupling; Tables S17–S18"],
            ["A2 GKT budget", "Disclosure done", "Abstract/Limitations/S21; GKT30 pooled Δ≈−0.038"],
            ["A3 cold-start Junyi", "Suppressed + GPU", "A3.5 bc9ea715; audit/2026-07-25-junyi-cold-stratum-nan-debug.md"],
            ["B4–B7 metrics/text", "Done", "|ρ| column; |C| vs |V_G|; edge-share; §3.2–3.7"],
            ["C8 page target", "34 pp main", "C8.4–C8.5: Fig→Supp S4–S7"],
            ["C9–C10", "Done", "refs scan; AI statement; Wilcoxon footnote"],
            ["PGS sign-off", "Pending", "Điều kiện nghiệm thu REVISION_PLAN §submit"],
        ],
    )

    add_heading(doc, "B. Mapping: simulated review ↔ PGS revision plan", 1)
    add_table(
        doc,
        ["Sim. comment", "PGS / plan item", "Status", "Notes"],
        [
            ["M1 GKT budget confound", "A2.0–A2.1", "Addressed (disclosure)", "Abstract qualifies observational; S21 exploratory; optional A2.3 GPU"],
            ["M2 three folds", "Limitations", "Acknowledged", "Wilcoxon p_min=0.25 footnote; bootstrap CI primary"],
            ["M3 narrow refs", "C9", "Partial", "32 refs + scholar scan; could add 5–8 KT baseline cites"],
            ["M4 beyond KT", "Future work", "Open", "1 paragraph sketch possible; not blocking GS"],
            ["M5 practitioner map", "Intro + Table 8", "Partial", "tab:audit-actions exists; boxed summary optional"],
            ["Injection mismatch", "A1", "Closed", "GPU verified; crossref clean/inject arms match cache"],
            ["Junyi very_cold artefact", "A3", "Closed", "n=11 fold0; models no longer identical AUC=1.0"],
            ["ECR/ρ scale", "B4", "Closed", "Column |ρ|; Eq.(8) aligned"],
            ["Page length", "C8", "Closed", "41→34 main pages"],
        ],
    )

    add_heading(doc, "1. Summary of the Submission", 1)
    add_para(
        doc,
        "The manuscript presents a leakage-controlled audit protocol for constructing and validating concept "
        "(knowledge-component, KC) graphs used by graph-enhanced knowledge tracing (KT) models. The authors argue "
        "that inferring prerequisite/similarity edges from the full interaction log before learner-disjoint splitting "
        "can let held-out information reach the model through the graph channel even when sequence training is "
        "fold-clean. The paper does not propose a new KT architecture; it contributes (C1) train-only fold-specific "
        "graph construction and DAG audit; (C2) DAG Disruption Rate (DDR) with prerequisite-preserving operator; "
        "(C3) KC-level cold-start stratification; (C4) ground-truth cross-validation on Junyi expert prerequisites; "
        "and (C5) an observational benchmarking note that GKT trails simpleKT by ≈0.041 AUC on XES3G5M under "
        "released (non-matched) budgets while GIKT remains competitive.",
    )
    add_para(
        doc,
        "Post-revision packaging: three public benchmarks + two synthetic corpora; three-fold learner CV; controlled "
        "leak injection (Section 4.3); multi-seed DDR→downstream sensitivity with manipulation check; appendix "
        "consolidates extended tables and Supplementary Figures S1–S7 (KT graphs, DDR curves, protocol plumbing, "
        "cold-start panels, GT PR curve). The authors frame the contribution as decision-support / reproducibility "
        "rather than headline accuracy gains (full-log vs train-only ablation ≤0.003 AUC).",
    )

    add_heading(doc, "2. Overall Assessment", 1)
    add_para(
        doc,
        "This remains a carefully engineered, unusually transparent methodology paper. After the July 2026 revision "
        "sprint, headline numbers were re-verified with an automated cross-reference script (12/12 checks, tolerance "
        "0.002 for fold means; exact match for injection arms vs GPU cache). The formal apparatus (leakage taxonomy, "
        "scalar diagnostics, DDR, cold-start strata, Algorithms 1–2) is precise; Limitations and two-factor scope "
        "paragraphs bound claims explicitly.",
    )
    add_para(
        doc,
        "Relative to a fresh external APIN review, residual weaknesses are narrower than in the pre-revision draft: "
        "(i) the GKT vs simpleKT gap is now budget-qualified throughout but still not compute-matched experimentally; "
        "(ii) reference breadth and APIN-general-audience positioning could still expand; (iii) three-fold CV limits "
        "significance resolution (acknowledged). For internal GS purposes, data-integrity and PGS-mandated fixes are "
        "addressed; remaining items are polish / optional experiments, not blockers.",
    )

    add_heading(doc, "3. Strengths", 1)
    for s in [
        "Clear, narrow, honestly scoped contribution; null full-log ablation (≤0.003 AUC) reported prominently.",
        "Rigorous formal apparatus with implementable algorithms and provenance artefacts.",
        "Manipulation-check design (p=0.90 graph destruction) makes DDR→AUC correlation conditionally interpretable.",
        "Controlled leak-injection demonstrates ECR-flag blind spot while builder throughput rises (Section 4.3).",
        "Numerical self-consistency: automated crossref + manual spot-checks on baseline, bootstrap, DDR-downstream, cold-start, leakage, DAG tables.",
        "Exemplary Limitations / two-factor scope-bounding; complete journal declarations including generative-AI disclosure.",
        "Revision traceability: REVISION_PLAN_APIN_NCS_20260722.md, audit memos, GPU runbooks, reproducible scripts.",
    ]:
        add_bullet(doc, s)

    add_heading(doc, "4. Major Comments (with revision status)", 1)

    add_heading(doc, "M1 — GKT vs simpleKT training-budget confound", 2)
    add_para(doc, "Original concern: −0.041 AUC headline under GKT 10 ep/batch 4 vs simpleKT 30 ep/batch 64.")
    add_para(doc, "Revision status: ADDRESSED (disclosure path). Abstract, Introduction, Limitations, and Table S21 now qualify the gap as observational; GKT30 nine-fold ablation (batch 32) narrows pooled Δ to ≈−0.038 without reversing sign. S21 changes four hyperparameters — appropriately flagged non-isolating.", bold=True)
    add_para(doc, "Remaining optional action: GPU compute-matched GKT/simpleKT (A2.3) only if PGS requires experimental closure.")

    add_heading(doc, "M2 — Three CV folds; underpowered significance tests", 2)
    add_para(doc, "Revision status: ACKNOWLEDGED. Wilcoxon p_min=0.25 footnote; paired-t bootstrap CIs carry primary inferential weight. Extending folds remains future work / optional.", bold=True)

    add_heading(doc, "M3 — Narrow reference list / related work", 2)
    add_para(doc, "Revision status: PARTIAL. C9 scholar scan completed; [6] Drive link removed; claims scoped vs [18]. Bibliography still 32 entries — acceptable for methodology focus; optional add 5–8 pyKT baseline citations.", bold=True)

    add_heading(doc, "M4 — Generalization beyond KT underdeveloped", 2)
    add_para(doc, "Revision status: OPEN (non-blocking). Section 6.2 / Future Work could add one concrete non-KT sketch; not required for GS sign-off.", bold=True)

    add_heading(doc, "M5 — Practitioner takeaway dispersed", 2)
    add_para(doc, "Revision status: PARTIAL. Four-step workflow in Introduction; decision map Table 8 (tab:audit-actions) in Discussion. Optional: boxed summary after abstract.", bold=True)

    add_heading(doc, "5. Minor Comments (updated)", 1)
    for s in [
        "Abstract density: partially trimmed in C8; GKT line now budget-qualified — further numeric trimming optional.",
        "Notation table: still recommended in appendix (not yet added).",
        "Main-text figures: by design, multi-benchmark KT/DDR/protocol/cold-start/GT panels moved to Supp S1–S7 for 34-page main body (C8.4–C8.5); XES3G5M claims rest on tables + S2.",
        "Table S21 / Future Work: epoch-and-batch-matched comparison already named as top follow-up in Limitations.",
        "Junyi GKT omission: compute-budget note present; optional wall-clock/GPU-hour sentence still useful.",
        "Software/hardware versions: consider adding PyTorch/CUDA/GPU model to Data availability paragraph.",
        "Cold-start terminology: defined in Section 3.4; optional one-line reminder in Discussion.",
    ]:
        add_bullet(doc, s)

    add_heading(doc, "6. Questions for the Authors (with prepared responses)", 1)
    qa = [
        (
            "Compute-matched GKT vs simpleKT on XES3G5M feasible?",
            "Primary path: disclosure + GKT30 ablation already in manuscript. Full match (A2.3) estimated ~1 week GPU — only if PGS mandates.",
        ),
        (
            "Natural high-throughput regime on public benchmarks?",
            "Authors' filter throttles throughput; injection + full-log ablation show low-throughput operating point on these corpora; two-factor table makes this explicit.",
        ),
        (
            "Extend DDR manipulation check to GIKT?",
            "Framed as future work; GKT/XES3G5M cell already passes check; DGEKT/ASSISTments graph-inert cells documented.",
        ),
    ]
    for q, a in qa:
        add_para(doc, f"Q: {q}", bold=True)
        add_para(doc, f"A: {a}")

    add_heading(doc, "7. Recommendation (internal use)", 1)
    add_para(
        doc,
        "For PGS advisor review: READY TO SEND. Simulated external stance: Minor Revision → after July sprint, "
        "most revision items are closed; remaining gaps are optional polish (refs, notation table, compute-matched GPU) "
        "rather than integrity or PGS-blocker issues. Recommend GS sign-off on current PDF + cover letter, then EM "
        "submission prep (metadata L05, clean MANIFEST, no .aux/.log in upload).",
    )

    add_heading(doc, "8. Verification Notes", 1)
    add_para(
        doc,
        "Automated: scripts/crossref_auc_numbers.py → audit/crossref_auc_report.md (25 July 2026, 0 failures). "
        "Manual pass: LaTeX source vs baseline_results.tex, bootstrap_ci_macros.tex, ddr_downstream_gkt.tex, "
        "cold_start_summary.tex, leakage_metrics.tex, dag_audit_summary.tex, graph_ablation.tex. "
        "Bibliography: 32 entries in refs_APIN.bib. Concerns above are design/framing (M1–M5), not fabricated numbers.",
    )

    add_heading(doc, "9. Suggested agenda for GS meeting (15–20 min)", 1)
    for s in [
        "Xác nhận framing A2 (−0.041 observational) — có cần A2.3 GPU compute-matched không?",
        "Duyệt nhanh bảng cold-start Junyi (suppressed) + injection decoupling narrative.",
        "Quyết định: mở rộng refs (M3) / notation table trước nộp EM hay để revision sau review journal?",
        "Ký nghiệm thu + gửi gói paper/gui_GS_package_APIN_20260725.zip.",
    ]:
        add_bullet(doc, s)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
