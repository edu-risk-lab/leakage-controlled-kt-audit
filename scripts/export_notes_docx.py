"""Export speaker notes from APIN_seminar_vi.pptx into note_seminar.docx."""
from pathlib import Path

from pptx import Presentation
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "paper" / "APIN_seminar_vi.pptx"
OUT = ROOT / "paper" / "note_seminar.docx"


def slide_title(slide) -> str:
    """Best-effort extraction of a slide's title text."""
    # Try the title placeholder first.
    try:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            t = slide.shapes.title.text.strip()
            if t:
                return t
    except Exception:
        pass
    # Fallback: first non-empty text shape.
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text.strip()
            if t:
                return t.split("\n")[0]
    return ""


def slide_note(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text.strip()


prs = Presentation(str(PPTX))
slides = list(prs.slides)
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(12)

title = doc.add_heading("Ghi chú thuyết trình — Seminar APIN", level=0)

sub = doc.add_paragraph(
    "Graph-Enhanced Knowledge Tracing: A Leakage-Audit Protocol"
)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic = True

doc.add_paragraph()

for idx, slide in enumerate(slides, start=1):
    title_text = slide_title(slide)
    note_text = slide_note(slide)

    head = doc.add_heading(level=1)
    run = head.add_run(f"Slide {idx}" + (f" — {title_text}" if title_text else ""))
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    if note_text:
        doc.add_paragraph(note_text)
    else:
        p = doc.add_paragraph("(Không có ghi chú)")
        p.runs[0].italic = True

doc.save(str(OUT))
print(f"Saved {OUT}  ({len(slides)} slides)")
